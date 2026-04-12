import secrets
from io import BytesIO
from fastapi import HTTPException, status
from openpyxl import load_workbook
from app.core.config import settings
from app.core.email_service import send_email
from app.core.security import create_access_token, hash_password, verify_password
from app.repositories.user_repository import UserRepository
from app.repositories.simple_repos import InstitutionRepository, DisciplineRepository, StudentRepository, EnrollmentRepository, VerificationTokenRepository, AttendanceEventRepository, AttendanceRequestRepository, AttendanceFinalRecordRepository
from app.schemas.models import *

class AuthService:
    def __init__(self):
        self.users=UserRepository(); self.students=StudentRepository(); self.tokens=VerificationTokenRepository()
    def signup_student(self, payload: SignupRequest):
        username=payload.username.strip().lower(); email=payload.email.strip().lower()
        if self.users.get_by_username(username): raise HTTPException(status_code=409, detail="Username already exists")
        if self.users.get_by_email(email): raise HTTPException(status_code=409, detail="Email already exists")
        user=self.users.create({"username": username, "password_hash": hash_password(payload.password), "full_name": payload.full_name.strip(), "email": email, "role": "student", "is_email_verified": False})
        self.students.create({"user_id": user["id"], "full_name": payload.full_name.strip(), "class_name": payload.class_name or "", "cpf": payload.cpf or "", "email": email, "phone": payload.phone or "", "created_by_user_id": user["id"]})
        token=secrets.token_urlsafe(32); self.tokens.create_token(user["id"], token, "email_verification")
        send_email(email, "Attend • Confirme seu email", f'<p>Confirme seu acesso: <a href="{settings.public_web_base_url}/verify-email?token={token}">validar email</a></p>')
        return user
    def login(self, payload: LoginRequest):
        user=self.users.get_by_username(payload.username.strip().lower())
        if not user or not verify_password(payload.password, user["password_hash"]): raise HTTPException(status_code=401, detail="Invalid credentials")
        return create_access_token(subject=user["id"], extra_claims={"role": user["role"], "username": user["username"]})
    def verify_email(self, token: str):
        rec=self.tokens.get_valid_token(token, "email_verification")
        if not rec: raise HTTPException(status_code=400, detail="Invalid or expired token")
        self.users.verify_email(rec["user_id"]); self.tokens.mark_used(rec["id"]); return {"status":"verified"}
    def forgot_password(self, payload: ForgotPasswordRequest):
        user=self.users.get_by_email(payload.email.strip().lower())
        if user:
            token=secrets.token_urlsafe(32); self.tokens.create_token(user["id"], token, "password_reset")
            send_email(user["email"], "Attend • Recuperação de senha", f'<p><a href="{settings.public_web_base_url}/reset-password?token={token}">redefinir senha</a></p>')
        return {"status":"ok"}
    def reset_password(self, payload: ResetPasswordRequest):
        rec=self.tokens.get_valid_token(payload.token, "password_reset")
        if not rec: raise HTTPException(status_code=400, detail="Invalid or expired token")
        self.users.update_password(rec["user_id"], hash_password(payload.new_password)); self.tokens.mark_used(rec["id"]); return {"status":"password_updated"}
    def first_access(self, payload: FirstAccessRequest):
        user=self.users.get_by_email(payload.email.strip().lower())
        if not user: raise HTTPException(status_code=404, detail="User not found")
        self.users.update_password(user["id"], hash_password(payload.new_password)); return {"status":"first_access_completed"}

class InstitutionService:
    def __init__(self): self.repo=InstitutionRepository()
    def create(self, payload: InstitutionCreate): return self.repo.create({**payload.model_dump(), "is_active": True})
    def list_all(self): return self.repo.list_all()
    def update(self, institution_id: str, payload: InstitutionCreate):
        if not self.repo.get_by_id(institution_id): raise HTTPException(status_code=404, detail="Institution not found")
        return self.repo.update(institution_id, payload.model_dump())

class DisciplineService:
    def __init__(self): self.repo=DisciplineRepository(); self.insts=InstitutionRepository()
    def create(self, payload: DisciplineCreate):
        if not self.insts.get_by_id(payload.institution_id): raise HTTPException(status_code=404, detail="Institution not found")
        return self.repo.create({**payload.model_dump(), "is_active": True})
    def list_all(self): return self.repo.list_all()

class StudentService:
    def __init__(self): self.users=UserRepository(); self.students=StudentRepository(); self.tokens=VerificationTokenRepository()
    def create(self, payload: StudentCreate, actor: dict):
        username=payload.username.strip().lower(); email=payload.email.strip().lower()
        if self.users.get_by_username(username): raise HTTPException(status_code=409, detail="Username already exists")
        if self.users.get_by_email(email): raise HTTPException(status_code=409, detail="Email already exists")
        temp_password=secrets.token_urlsafe(8)
        user=self.users.create({"username": username, "password_hash": hash_password(temp_password), "full_name": payload.full_name.strip(), "email": email, "role": "student", "is_email_verified": False})
        student=self.students.create({"user_id": user["id"], "full_name": payload.full_name.strip(), "class_name": payload.class_name.strip(), "cpf": payload.cpf.strip(), "email": email, "phone": payload.phone.strip(), "created_by_user_id": actor["id"]})
        if payload.send_email:
            token=secrets.token_urlsafe(32); self.tokens.create_token(user["id"], token, "email_verification")
            send_email(email, "Attend • Primeiro acesso", f'<p>Usuário: <strong>{username}</strong></p><p>Senha provisória: <strong>{temp_password}</strong></p><p><a href="{settings.public_web_base_url}/first-access?email={email}">primeiro acesso</a></p>')
        return {**student, "user_id": user["id"]}
    def list_all(self): return self.students.list_all()
    def update(self, student_id: str, payload: StudentUpdate):
        if not self.students.get_by_id(student_id): raise HTTPException(status_code=404, detail="Student not found")
        return self.students.update(student_id, payload.model_dump())
    def import_excel(self, file_bytes: bytes, actor: dict):
        wb=load_workbook(filename=BytesIO(file_bytes)); ws=wb.active
        headers=[str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        expected=["nome_completo","turma","cpf","email","telefone","username","enviar_email"]
        if headers!=expected: raise HTTPException(status_code=400, detail=f"Excel headers must be exactly: {expected}")
        created=0; updated=0; errors=[]
        for idx,row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row): continue
            try:
                full_name,class_name,cpf,email,phone,username,send_email=row
                existing=self.users.get_by_email(str(email).strip().lower())
                if existing:
                    student=self.students.get_by_user_id(existing["id"])
                    if student:
                        self.students.update(student["id"], {"full_name": str(full_name).strip(), "class_name": str(class_name).strip(), "cpf": str(cpf).strip(), "email": str(email).strip().lower(), "phone": str(phone).strip()})
                        updated += 1; continue
                self.create(StudentCreate(full_name=str(full_name).strip(), class_name=str(class_name).strip(), cpf=str(cpf).strip(), email=str(email).strip().lower(), phone=str(phone).strip(), username=str(username).strip().lower(), send_email=str(send_email).strip().lower()=="true"), actor)
                created += 1
            except Exception as exc:
                errors.append(f"linha {idx}: {exc}")
        return {"created": created, "updated": updated, "errors": errors}

class EnrollmentService:
    def __init__(self): self.repo=EnrollmentRepository(); self.users=UserRepository(); self.disc=DisciplineRepository()
    def create(self, payload: EnrollmentCreate):
        user=self.users.get_by_id(payload.student_user_id)
        if not user or user["role"]!="student": raise HTTPException(status_code=400, detail="Student not found")
        if not self.disc.get_by_id(payload.discipline_id): raise HTTPException(status_code=404, detail="Discipline not found")
        if self.repo.find_active(payload.student_user_id, payload.discipline_id): raise HTTPException(status_code=409, detail="Enrollment already exists")
        return self.repo.create(payload.model_dump())
    def list_all(self): return self.repo.list_all()

class AttendanceService:
    def __init__(self):
        self.events=AttendanceEventRepository(); self.requests=AttendanceRequestRepository(); self.records=AttendanceFinalRecordRepository(); self.disc=DisciplineRepository(); self.enroll=EnrollmentRepository()
    def create_event(self, payload: AttendanceEventCreate, actor: dict):
        if not self.disc.get_by_id(payload.discipline_id): raise HTTPException(status_code=404, detail="Discipline not found")
        return self.events.create({"discipline_id": payload.discipline_id, "title": payload.title.strip(), "lesson_date": payload.lesson_date.strip(), "lesson_time": payload.lesson_time.strip(), "teacher_user_id": actor["id"], "allow_remote_requests": payload.allow_remote_requests})
    def list_events(self): return self.events.list_all()
    def create_request(self, payload: AttendanceRequestCreate, actor: dict):
        event=self.events.get_by_id(payload.event_id)
        if not event: raise HTTPException(status_code=404, detail="Event not found")
        if event["status"]!="open": raise HTTPException(status_code=400, detail="Event is closed")
        if not self.enroll.find_active(actor["id"], event["discipline_id"]): raise HTTPException(status_code=400, detail="Student is not enrolled in this discipline")
        if self.requests.find_existing_pending(payload.event_id, actor["id"]): raise HTTPException(status_code=409, detail="Pending request already exists")
        return self.requests.create({"event_id": payload.event_id, "student_user_id": actor["id"], "channel": payload.channel, "stars": payload.stars, "tags": payload.tags, "comments": payload.comments})
    def list_requests(self, event_id: str): return self.requests.list_by_event(event_id)
    def review_request(self, request_id: str, status_value: str, actor: dict):
        req=self.requests.get_by_id(request_id)
        if not req: raise HTTPException(status_code=404, detail="Request not found")
        event=self.events.get_by_id(req["event_id"])
        self.requests.review(request_id, status_value, actor["id"])
        final_status="confirmed" if status_value=="confirmed" else "unconfirmed"
        return self.records.upsert(event["id"], req["student_user_id"], {"discipline_id": event["discipline_id"], "status": final_status, "justification": None})
    def close_event(self, event_id: str):
        event=self.events.get_by_id(event_id)
        if not event: raise HTTPException(status_code=404, detail="Event not found")
        existing={r["student_user_id"] for r in self.records.list_by_event(event_id)}
        for enr in self.enroll.list_by_discipline(event["discipline_id"]):
            if enr["student_user_id"] not in existing:
                self.records.upsert(event_id, enr["student_user_id"], {"discipline_id": event["discipline_id"], "status": "absent", "justification": None})
        self.events.close(event_id); return self.events.get_by_id(event_id)
    def list_final_records(self, event_id: str): return self.records.list_by_event(event_id)
